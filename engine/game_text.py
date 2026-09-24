#!/usr/bin/env python3
"""
Read the game's own text and the framework's numbers from 'HEARsay game text.xlsx'.

    python3 game_text.py ["HEARsay game text.xlsx"]

read_game_text() returns (game_text, errors). game_text.text maps each key to its text, game_text.facts maps each
fact key to its value, and game_text.rows / .passages / .fact_rows keep the full rows for check_framework.py.
Needs openpyxl (pip3 install openpyxl).
"""
import os, re, sys
from types import SimpleNamespace

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
from scenario_matrix import _sheet, _keyed

try:
    from openpyxl import load_workbook
except ImportError:
    sys.exit("This needs openpyxl: pip3 install openpyxl")

DEFAULT = os.path.join(os.path.dirname(os.path.abspath(__file__)), "HEARsay game text.xlsx")
# Placeholders every text can use; the engine fills them from the scenario and the facts.
GLOBAL_PLACEHOLDERS = {"budget", "n_witnesses", "n_questions", "n_major", "bonus", "rank_exact", "rank_near", "ccc", "title"}
CHECKS = {"verbatim", "paraphrase", "label"}
# Facts the build and the engine rely on.
REQUIRED_FACTS = ["witnesses", "questions", "question_budget", "tripwires_per_witness", "clues_per_witness",
                  "red_herrings_per_witness", "clues", "givers_per_clue", "red_herrings", "givers_per_red_herring",
                  "rh_set_min", "rh_set_max", "major_clues", "clues_per_major", "followups_per_witness", "submissions",
                  "followup_bonus", "rank_exact", "rank_near"]


def read_game_text(path=DEFAULT):
    errors = []
    wb = load_workbook(path, data_only=True)

    passages = _sheet(wb, "Passages", ["Engine passage", "Framework name", "Framework section", "Notes"], errors)
    engine_passages = {p["Engine passage"] for p in passages}

    rows = _sheet(wb, "Text", ["Key", "Engine passage", "Text", "Framework section", "Check", "Placeholders"], errors)
    text = {}
    for r in rows:
        k, where = r["Key"], f"Text row {r['_row']} ({r['Key'] or 'no key'})"
        if not k:
            errors.append(f"Text row {r['_row']}: Key is blank"); continue
        if k in text:
            errors.append(f"{where}: key appears twice"); continue
        if r["Engine passage"] not in engine_passages:
            errors.append(f"{where}: engine passage {r['Engine passage']!r} is not on the Passages sheet")
        check = r["Check"].lower()
        if check not in CHECKS:
            errors.append(f"{where}: Check must be verbatim, paraphrase or label, not {r['Check']!r}")
        if check in ("verbatim", "paraphrase") and not r["Framework section"]:
            errors.append(f"{where}: a {check} row needs a Framework section")
        allowed = GLOBAL_PLACEHOLDERS | set(r["Placeholders"].split())
        for ph in re.findall(r"\{(\w+)\}", r["Text"]):
            if ph not in allowed:
                errors.append(f"{where}: {{{ph}}} is not a placeholder this text can use "
                              f"(allowed: {', '.join(sorted(allowed))})")
        if check == "verbatim" and "{" in r["Text"]:
            errors.append(f"{where}: a verbatim row cannot contain placeholders")
        text[k] = r["Text"]

    fact_rows = _keyed(_sheet(wb, "Facts", ["Key", "Fact", "Value", "Framework wording", "Framework section"], errors), "Key", "Facts", errors)
    facts = {}
    for k, r in fact_rows.items():
        try:
            v = float(r["Value"])
            facts[k] = int(v) if v.is_integer() else v
        except ValueError:
            errors.append(f"Facts {k}: Value must be a number, not {r['Value']!r}")
    for k in REQUIRED_FACTS:
        if k not in facts:
            errors.append(f"Facts: '{k}' is missing")
    return SimpleNamespace(text=text, facts=facts, rows=rows, passages=passages, fact_rows=list(fact_rows.values())), errors


def main():
    path = sys.argv[1] if len(sys.argv) > 1 else DEFAULT
    g, errors = read_game_text(path)
    print(f"{len(g.text)} texts, {len(g.passages)} passages, {len(g.facts)} facts")
    print("PROBLEMS:" if errors else "OK: the game text passes every check")
    for e in errors:
        print(" -", e)
    sys.exit(1 if errors else 0)


if __name__ == "__main__":
    main()
