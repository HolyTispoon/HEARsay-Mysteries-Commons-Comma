#!/usr/bin/env python3
"""
Read a HEARsay scenario from its Response Matrix workbook (.xlsx) and check it.

    python3 scenario_matrix.py "SCENARIO – Response Matrix.xlsx"

The workbook is the source the game is built from. read_matrix() returns the scenario plus a list of
problems: missing fields, broken references, the framework's structural rules (tripwires, clue counts,
receptiveness ranks, with the numbers taken from the Facts sheet of 'HEARsay game text.xlsx'), and places
where the workbook disagrees with itself (the Grid against the Long table, 'Given by' columns, repeated set
titles). Scenario text is filed under the framework's passage names (Passages sheet). Columns are found by
their header, so their order does not matter.

Needs openpyxl (pip3 install openpyxl).
"""
import re, sys
from collections import Counter
from types import SimpleNamespace

try:
    from openpyxl import load_workbook
except ImportError:
    sys.exit("This needs openpyxl: pip3 install openpyxl")

ENDING_KEYS = ("correct", "right_person_wrong_place", "wrong", "out_of_questions")
# Scenario text on the Passages sheet: (framework passage, part) -> (where it goes, required?)
PASSAGE_PARTS = {
    ("Start", "Summary"): ("pitch", False),
    ("the Hook", "Situation"): ("player_role", True),
    ("the Hook", "Assignment"): ("briefing", True),
    ("the Hook", "CCC"): ("ccc_name", True),
    ("the Hook", "CCC summary"): ("ccc_summary", True),
    ("the Hook", "Clause"): ("clause", False),
    ("the Hook", "Clause explainer"): ("clause_explainer", False),
    ("Submission", "Out of questions"): ("out_of_questions", True),
    ("ending", "Correct"): ("correct", True),
    ("ending", "Right person, wrong place"): ("right_person_wrong_place", True),
    ("ending", "Wrong"): ("wrong", True),
}


def _text(v):
    if v is None:
        return ""
    if isinstance(v, float) and v.is_integer():
        v = int(v)
    return str(v).strip()


def _sheet(wb, name, headers, errors):
    """Rows of a sheet as dicts keyed by header, with their spreadsheet row numbers. Blank rows are skipped."""
    if name not in wb.sheetnames:
        errors.append(f"Missing sheet '{name}'")
        return []
    ws = wb[name]
    head = [_text(c.value) for c in ws[1]]
    missing = [h for h in headers if h not in head]
    if missing:
        errors.append(f"{name}: missing column(s) {', '.join(repr(m) for m in missing)}")
        return []
    col = {h: head.index(h) for h in headers}
    rows = []
    for r, values in enumerate(ws.iter_rows(min_row=2, values_only=True), start=2):
        row = {h: _text(values[i]) if i < len(values) else "" for h, i in col.items()}
        if any(row.values()):
            row["_row"] = r
            rows.append(row)
    return rows


def _keyed(rows, key, sheet, errors):
    """Rows as a dict keyed by one column; flags blanks and duplicates."""
    out = {}
    for r in rows:
        k = r[key]
        if not k:
            errors.append(f"{sheet} row {r['_row']}: '{key}' is blank")
        elif k in out:
            errors.append(f"{sheet} row {r['_row']}: {key} {k} appears twice")
        else:
            out[k] = r
    return out


def _ids(s):
    return [x for x in re.split(r"[,\s;]+", s) if x and x.lower() != "none"]


def read_matrix(path, facts=None):
    """Return (scenario, errors). facts: the framework numbers (default: read from 'HEARsay game text.xlsx')."""
    errors = []
    if facts is None:
        from game_text import read_game_text
        g, gerr = read_game_text()
        if gerr:
            errors.append(f"HEARsay game text.xlsx has problems ({len(gerr)}; run game_text.py)")
        facts = g.facts
    wb = load_workbook(path, data_only=True)

    def need(value, where):
        if not value:
            errors.append(f"{where} is blank")
        return value

    # Settings
    settings = {r["Setting"]: r["Value"] for r in _sheet(wb, "Settings", ["Setting", "Value"], errors)}
    def setting(k, required=True):
        v = settings.get(k, "")
        if required and not v:
            errors.append(f"Settings: '{k}' is missing or blank")
        return v
    # Passages: scenario text filed under the framework's passage names
    passages = {}
    for r in _sheet(wb, "Passages", ["Passage", "Part", "Text"], errors):
        key = next((k for k in PASSAGE_PARTS if k[0].lower() == r["Passage"].lower() and k[1].lower() == r["Part"].lower()), None)
        if not key:
            errors.append(f"Passages row {r['_row']}: unknown passage/part {r['Passage']!r} / {r['Part']!r} "
                          f"(expected one of: {'; '.join(f'{a} / {b}' for a, b in PASSAGE_PARTS)})")
        elif PASSAGE_PARTS[key][0] in passages:
            errors.append(f"Passages row {r['_row']}: {key[0]} / {key[1]} appears twice")
        else:
            passages[PASSAGE_PARTS[key][0]] = r["Text"]
    for (pas, part), (field, required) in PASSAGE_PARTS.items():
        if required and not passages.get(field):
            errors.append(f"Passages: {pas} / {part} is missing or blank")

    # Mystery
    mystery = {r["Field"]: r["Text"] for r in _sheet(wb, "Mystery", ["Field", "Text"], errors)}
    def my(k):
        return need(mystery.get(k, ""), f"Mystery: '{k}'")

    # Final-answer form
    who, where, correct = [], [], {"Person": [], "Place": []}
    for r in _sheet(wb, "Final-answer form", ["Kind", "Option", "Correct"], errors):
        kind = r["Kind"].capitalize()
        if kind not in correct:
            errors.append(f"Final-answer form row {r['_row']}: Kind must be Person or Place, not {r['Kind']!r}")
            continue
        (who if kind == "Person" else where).append(need(r["Option"], f"Final-answer form row {r['_row']}: Option"))
        if r["Correct"].lower() in ("yes", "y", "x", "true", "correct"):
            correct[kind].append(r["Option"])
    for kind, found in correct.items():
        if len(found) != 1:
            errors.append(f"Final-answer form: exactly one {kind} must be marked correct (found {len(found)})")

    # Questions
    qrows = _keyed(_sheet(wb, "Questions", ["ID", "Question", "Angle", "Tripwire for", "Why it lands"], errors), "ID", "Questions", errors)
    questions = [{"id": k, "text": need(r["Question"], f"Questions {k}: Question"), "angle": r["Angle"],
                  "tripwire_for_name": r["Tripwire for"], "tripwire_logic": r["Why it lands"]} for k, r in qrows.items()]
    Q = {q["id"]: q for q in questions}

    # Witnesses
    wcols = ["ID", "Name", "Role", "Public background card", "Position on the CCC", "Withheld secret", "Tripwire", "Tripwire question",
             "Walk-away text", "Receptiveness rank", "Receptive on", "Location", "Portrait"]
    wrows = _keyed(_sheet(wb, "Witnesses", wcols, errors), "ID", "Witnesses", errors)
    witnesses, tripwire_responses = [], {}
    for k, r in wrows.items():
        where_ = f"Witnesses {k}"
        try:
            rank = int(float(r["Receptiveness rank"]))
        except ValueError:
            errors.append(f"{where_}: Receptiveness rank must be a number, not {r['Receptiveness rank']!r}")
            rank = 0
        witnesses.append({
            "id": k, "name": need(r["Name"], f"{where_}: Name"), "portrait": need(r["Portrait"], f"{where_}: Portrait"),
            "role": need(r["Role"], f"{where_}: Role"), "public_card": need(r["Public background card"], f"{where_}: Public background card"),
            "stance": r["Position on the CCC"], "secret": r["Withheld secret"], "tripwire": need(r["Tripwire"], f"{where_}: Tripwire"),
            "receptiveness_rank": rank, "receptive_questions": _ids(r["Receptive on"]),
            "location": need(r["Location"], f"{where_}: Location")})
        tripwire_responses[k] = need(r["Walk-away text"], f"{where_}: Walk-away text")
        if r["Tripwire"] in Q and r["Tripwire question"] != Q[r["Tripwire"]]["text"]:
            errors.append(f"{where_}: 'Tripwire question' does not match the text of {r['Tripwire']} on the Questions sheet")
    W = {w["id"]: w for w in witnesses}
    by_name = {w["name"]: w["id"] for w in witnesses}
    for q in questions:
        q["tripwire_for"] = by_name.get(q.pop("tripwire_for_name"), "")
        if not q["tripwire_for"]:
            errors.append(f"Questions {q['id']}: 'Tripwire for' must be a witness name from the Witnesses sheet")

    # Clues, Major clues, red herrings
    crows = _keyed(_sheet(wb, "Clues", ["ID", "Major clue group", "Clue", "Given by"], errors), "ID", "Clues", errors)
    clues = {k: {"group": r["Major clue group"], "text": need(r["Clue"], f"Clues {k}: Clue")} for k, r in crows.items()}
    mrows = _keyed(_sheet(wb, "Major Clues", ["ID", "Letter", "Title", "Clues", "Text shown to player"], errors), "ID", "Major Clues", errors)
    major = [{"id": k, "letter": need(r["Letter"], f"Major Clues {k}: Letter"), "title": need(r["Title"], f"Major Clues {k}: Title"),
              "clues": _ids(r["Clues"]), "text": need(r["Text shown to player"], f"Major Clues {k}: Text shown to player")}
             for k, r in mrows.items()]
    rrows = _keyed(_sheet(wb, "Red Herrings", ["ID", "Set", "Set title", "Red Herring", "Given by", "How the set is discounted"], errors),
                   "ID", "Red Herrings", errors)
    red_herrings, rh_sets = {}, {}
    for k, r in rrows.items():
        red_herrings[k] = {"set": need(r["Set"], f"Red Herrings {k}: Set"), "text": need(r["Red Herring"], f"Red Herrings {k}: Red Herring")}
        s = {"title": r["Set title"], "resolution": r["How the set is discounted"]}
        if r["Set"] in rh_sets and rh_sets[r["Set"]] != s:
            errors.append(f"Red Herrings {k}: set {r['Set']} has a different title or resolution from its earlier rows")
        rh_sets.setdefault(r["Set"], s)

    # Long table: the answers
    lcols = ["Witness ID", "Witness", "Question ID", "Question", "Item", "Item type", "Group / set", "Opinion receptive", "Answer text"]
    responses = {w["id"]: {} for w in witnesses}
    receptive_cells = {w["id"]: set() for w in witnesses}
    cells = {}
    for r in _sheet(wb, "Long table", lcols, errors):
        wid, qid, where_ = r["Witness ID"], r["Question ID"], f"Long table row {r['_row']}"
        if wid not in W or qid not in Q:
            errors.append(f"{where_}: unknown witness {wid!r} or question {qid!r}")
            continue
        if (wid, qid) in cells:
            errors.append(f"{where_}: {wid} × {qid} appears twice")
            continue
        cells[(wid, qid)] = r
        if r["Witness"] != W[wid]["name"] or r["Question"] != Q[qid]["text"]:
            errors.append(f"{where_}: the Witness or Question column does not match the Witnesses/Questions sheets")
        item = r["Item"]
        if W[wid]["tripwire"] == qid:
            if item.upper() != "TRIPWIRE":
                errors.append(f"{where_}: {qid} is {W[wid]['name']}'s tripwire, so Item must be TRIPWIRE")
            elif r["Answer text"] != tripwire_responses[wid]:
                errors.append(f"{where_}: tripwire text differs from the Walk-away text on the Witnesses sheet")
            continue
        if item in clues:
            kind, group = "clue", clues[item]["group"]
        elif item in red_herrings:
            kind, group = "red herring", red_herrings[item]["set"]
        else:
            errors.append(f"{where_}: Item {item!r} is not a clue or red herring ID")
            continue
        if r["Item type"] != kind or r["Group / set"] != group:
            errors.append(f"{where_}: Item type / Group should be '{kind}' / '{group}' for {item}")
        flag = r["Opinion receptive"].lower()
        if flag not in ("yes", "no"):
            errors.append(f"{where_}: Opinion receptive must be yes or no")
        elif flag == "yes":
            receptive_cells[wid].add(qid)
        responses[wid][qid] = {"item": item, "text": need(r["Answer text"], f"{where_}: Answer text")}
    for w in witnesses:
        if set(w["receptive_questions"]) != receptive_cells[w["id"]]:
            errors.append(f"Witnesses {w['id']}: 'Receptive on' ({', '.join(w['receptive_questions']) or 'none'}) does not match "
                          f"the Long table ({', '.join(sorted(receptive_cells[w['id']])) or 'none'})")

    # Given-by columns
    givers = {}
    for w in witnesses:
        for q in questions:
            a = responses[w["id"]].get(q["id"])
            if a:
                givers.setdefault(a["item"], []).append(f"{w['name']} ({q['id']})")
    for k, r in list(crows.items()) + list(rrows.items()):
        sheet = "Clues" if k in crows else "Red Herrings"
        if sorted(_ids_given(r["Given by"])) != sorted(givers.get(k, [])):
            errors.append(f"{sheet} {k}: 'Given by' says {r['Given by']!r}, but the Long table has {'; '.join(givers.get(k, [])) or 'nobody'}")

    # Grid (a checked copy of the Long table)
    if "Grid" in wb.sheetnames:
        ws = wb["Grid"]
        head = [_text(c.value) for c in ws[1]]
        for r, values in enumerate(ws.iter_rows(min_row=2, values_only=True), start=2):
            m = re.match(r"(Q\d+)\.", _text(values[0]))
            if not m:
                continue
            qid = m.group(1)
            for i, h in enumerate(head[1:], start=1):
                wid = by_name.get(h)
                cell = cells.get((wid, qid))
                if not cell:
                    continue
                got = _text(values[i]) if i < len(values) else ""
                if W[wid]["tripwire"] == qid:
                    want = "TW — " + cell["Answer text"]
                else:
                    flag = "receptive" if cell["Opinion receptive"].lower() == "yes" else "unreceptive"
                    want = f"[{cell['Item']} | {flag}] {cell['Answer text']}"
                if got != want:
                    errors.append(f"Grid {qid} × {h}: does not match the Long table (the game reads the Long table)")

    # Follow-ups and endings
    follow_ups = {}
    for r in _sheet(wb, "Follow-ups", ["Witness ID", "Receptive option", "Unreceptive option", "Explanation on success", "Walk-away on failure"], errors):
        where_ = f"Follow-ups {r['Witness ID']}"
        follow_ups[r["Witness ID"]] = {"good": need(r["Receptive option"], f"{where_}: Receptive option"),
                                       "bad": need(r["Unreceptive option"], f"{where_}: Unreceptive option"),
                                       "why": need(r["Explanation on success"], f"{where_}: Explanation on success"),
                                       "walkout": need(r["Walk-away on failure"], f"{where_}: Walk-away on failure")}
    s = SimpleNamespace(
        TITLE=setting("Title"), SCENARIO_ID=setting("Scenario ID"), IFID=setting("IFID"),
        FACTS=facts, QUESTION_BUDGET=facts.get("question_budget", 0), FOLLOWUP_BONUS=facts.get("followup_bonus", 0),
        MAP=setting("Map", False) or "campus-map.png", PLAY_URL=setting("Play URL", False), PITCH=passages.get("pitch", ""),
        PUZZLE={"mystery": my("The mystery"), "solution": my("The solution"),
                "who_label": my("Person question"), "where_label": my("Place question"),
                "final_answer": {"who": (correct["Person"] or [""])[0], "where": (correct["Place"] or [""])[0]},
                "who_options": who, "where_options": where},
        CONTROVERSY={"name": passages.get("ccc_name", ""), "summary": passages.get("ccc_summary", ""),
                     "clause": passages.get("clause", ""), "clause_explainer": passages.get("clause_explainer", "")},
        SETTING={"player_role": passages.get("player_role", ""), "briefing": passages.get("briefing", "")},
        QUESTIONS=questions, WITNESSES=witnesses, CLUES=clues, MAJOR_CLUES=major,
        RED_HERRINGS=red_herrings, RED_HERRING_SETS=rh_sets, FOLLOW_UPS=follow_ups,
        TRIPWIRE_RESPONSES=tripwire_responses, RESPONSES=responses,
        ENDINGS={k: passages.get(k, "") for k in ENDING_KEYS},
        BONUS_RANKING={"answer": [w["id"] for w in sorted(witnesses, key=lambda w: w["receptiveness_rank"])]},
    )
    if facts:
        errors += structural_errors(s, facts)
    return s, errors


def _ids_given(s):
    return [x.strip() for x in s.split(";") if x.strip()]


def structural_errors(d, f):
    """The framework's structural rules, with the numbers from the Facts sheet (f)."""
    errors = []
    N_WITNESSES, N_QUESTIONS, N_CLUES, N_RED_HERRINGS, N_MAJOR = f["witnesses"], f["questions"], f["clues"], f["red_herrings"], f["major_clues"]
    CLUES_PER_WITNESS, RH_PER_WITNESS, CLUES_PER_MAJOR = f["clues_per_witness"], f["red_herrings_per_witness"], f["clues_per_major"]
    qids = [q["id"] for q in d.QUESTIONS]
    if len(d.WITNESSES) != N_WITNESSES: errors.append(f"need {N_WITNESSES} witnesses, found {len(d.WITNESSES)}")
    if len(qids) != N_QUESTIONS: errors.append(f"need {N_QUESTIONS} questions, found {len(qids)}")
    tw = [w["tripwire"] for w in d.WITNESSES]
    if sorted(tw) != sorted(qids):
        errors.append(f"tripwires must be one per question: {', '.join(tw)}")
    for q in d.QUESTIONS:
        owner = next((w for w in d.WITNESSES if w["tripwire"] == q["id"]), None)
        if owner and owner["id"] != q["tripwire_for"]:
            errors.append(f"{q['id']}: Questions says it is {q['tripwire_for'] or '?'}'s tripwire, Witnesses says {owner['id']}'s")

    clue_count, rh_count = Counter(), Counter()
    for w in d.WITNESSES:
        r = d.RESPONSES.get(w["id"], {})
        expected = sorted(q for q in qids if q != w["tripwire"])
        if sorted(r) != expected:
            errors.append(f"{w['id']} answers {', '.join(sorted(r))}; expected {', '.join(expected)}")
        items = [a["item"] for a in r.values()]
        cl = [i for i in items if i in d.CLUES]
        rh = [i for i in items if i in d.RED_HERRINGS]
        if len(cl) != CLUES_PER_WITNESS or len(rh) != RH_PER_WITNESS:
            errors.append(f"{w['id']}: {len(cl)} clues / {len(rh)} red herrings (need {CLUES_PER_WITNESS} / {RH_PER_WITNESS})")
        if len(set(cl)) != len(cl):
            errors.append(f"{w['id']}: gives the same clue twice")
        clue_count.update(cl); rh_count.update(rh)
        want = N_QUESTIONS - w["receptiveness_rank"]
        if len(w["receptive_questions"]) != want:
            errors.append(f"{w['id']}: rank {w['receptiveness_rank']} needs {want} receptive answers, has {len(w['receptive_questions'])}")
        if w["tripwire"] in w["receptive_questions"]:
            errors.append(f"{w['id']}: tripwire listed as receptive")
    ranks = sorted(w["receptiveness_rank"] for w in d.WITNESSES)
    if ranks != list(range(1, len(d.WITNESSES) + 1)):
        errors.append(f"receptiveness ranks must be 1 to {len(d.WITNESSES)}, each once: {ranks}")

    for c in d.CLUES:
        if clue_count[c] != f["givers_per_clue"]: errors.append(f"{c} is given {clue_count[c]} times (need {f['givers_per_clue']})")
    for r in d.RED_HERRINGS:
        if rh_count[r] != f["givers_per_red_herring"]: errors.append(f"{r} is given {rh_count[r]} times (need {f['givers_per_red_herring']})")
    if len(d.CLUES) != N_CLUES: errors.append(f"need {N_CLUES} clues, found {len(d.CLUES)}")
    if len(d.RED_HERRINGS) != N_RED_HERRINGS: errors.append(f"need {N_RED_HERRINGS} red herrings, found {len(d.RED_HERRINGS)}")
    if len(d.MAJOR_CLUES) != N_MAJOR: errors.append(f"need {N_MAJOR} Major Clues, found {len(d.MAJOR_CLUES)}")
    mc = [c for m in d.MAJOR_CLUES for c in m["clues"]]
    if sorted(mc) != sorted(d.CLUES): errors.append("the Major Clues must each list five clues, using every clue once")
    for m in d.MAJOR_CLUES:
        if len(m["clues"]) != CLUES_PER_MAJOR: errors.append(f"{m['id']} lists {len(m['clues'])} clues (need {CLUES_PER_MAJOR})")
        for c in m["clues"]:
            if c in d.CLUES and d.CLUES[c]["group"] != m["id"]:
                errors.append(f"{c}: Clues sheet puts it in {d.CLUES[c]['group']}, Major clues puts it in {m['id']}")
    sets = Counter(r["set"] for r in d.RED_HERRINGS.values())
    for s, n in sets.items():
        if not f["rh_set_min"] <= n <= f["rh_set_max"]:
            errors.append(f"Red Herring set {s} has {n} members (the framework allows {f['rh_set_min']} to {f['rh_set_max']})")
    for w in d.WITNESSES:
        if w["id"] not in d.FOLLOW_UPS: errors.append(f"no follow-up for {w['id']}")
    for k in d.FOLLOW_UPS:
        if k not in {w["id"] for w in d.WITNESSES}: errors.append(f"Follow-ups: unknown witness {k}")
    return errors


def main():
    if len(sys.argv) != 2:
        sys.exit(__doc__)
    s, errors = read_matrix(sys.argv[1])
    n = sum(len(w["receptive_questions"]) for w in s.WITNESSES)
    print(f"{s.TITLE}: {len(s.WITNESSES)} witnesses, {len(s.CLUES)} clues, {len(s.RED_HERRINGS)} red herrings, "
          f"{n} receptive answers, budget {s.QUESTION_BUDGET}")
    print("PROBLEMS:" if errors else "OK: the matrix passes every check")
    for e in errors:
        print(" -", e)
    sys.exit(1 if errors else 0)


if __name__ == "__main__":
    main()
